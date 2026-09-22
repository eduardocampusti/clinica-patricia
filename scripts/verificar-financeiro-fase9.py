from __future__ import annotations

import html
import re
import warnings
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

warnings.filterwarnings("ignore", message="Workbook contains no default style")


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "scratch" / "fase9-relatorios"
PDF = OUT / "RELATORIO_FINANCEIRO_EXEMPLO_SINTETICO.pdf"
XLSX = OUT / "RELATORIO_FINANCEIRO_EXEMPLO_SINTETICO.xlsx"
MALICIOUS = (
    "=1+1",
    "+SUM(A1:A2)",
    "-1+1",
    "@TESTE",
    "+CMD|' /C calc'!A0",
    "@SUM(A1:A2)",
    '<script>alert(1)</script>',
    '../CLINICA/"TESTE"',
    '=HYPERLINK("https://invalid.example")',
)


def verify_pdf() -> tuple[int, int]:
    reader = PdfReader(PDF)
    root = reader.trailer["/Root"]
    assert "/AcroForm" not in root, "PDF contém formulário ativo"
    open_action = root.get("/OpenAction")
    if open_action is not None:
        assert isinstance(open_action, list), "PDF contém ação automática executável"
        assert len(open_action) >= 2 and str(open_action[1]) in ("/Fit", "/FitH", "/FitV"), "PDF contém ação automática não permitida"
    names = root.get("/Names")
    if names:
        assert "/JavaScript" not in names, "PDF contém JavaScript"
    for page in reader.pages:
        assert "/Annots" not in page, "PDF contém anotações ou links ativos"
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    for payload in MALICIOUS[:-1]:
        assert payload in text, f"Texto hostil não permaneceu literal no PDF: {payload}"
    assert not re.search(r"\b[0-9a-f]{8}-[0-9a-f]{4}-[1-5][0-9a-f]{3}-[89ab][0-9a-f]{3}-[0-9a-f]{12}\b", text, re.I)
    assert not re.search(r"\b\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b", text)
    return len(reader.pages), len(text)


def verify_xlsx() -> tuple[int, int]:
    with zipfile.ZipFile(XLSX) as archive:
        xml_entries = [name for name in archive.namelist() if name.endswith(".xml")]
        xml = "\n".join(archive.read(name).decode("utf-8") for name in xml_entries)
    assert not re.search(r"<f(?:\s|>)", xml), "XLSX contém fórmula executável"
    decoded = html.unescape(xml)
    text_nodes = "\n".join(html.unescape(value) for value in re.findall(r"<t(?:\s[^>]*)?>(.*?)</t>", xml, re.S))
    for payload in MALICIOUS:
        assert payload in decoded, f"Texto hostil não permaneceu literal no XLSX: {payload}"
    assert "DDE" not in xml.upper()
    assert not re.search(r"\b[0-9a-f]{8}-[0-9a-f]{4}-[1-5][0-9a-f]{3}-[89ab][0-9a-f]{3}-[0-9a-f]{12}\b", text_nodes, re.I)
    assert not re.search(r"\b\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b", text_nodes)

    workbook = load_workbook(XLSX, read_only=True, data_only=False)
    assert workbook.sheetnames == ["Resumo", "Recebimentos", "Repasses", "Fiscal"]
    visible_text = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f", f"Fórmula encontrada em {sheet.title}!{cell.coordinate}"
                if isinstance(cell.value, str):
                    visible_text.append(cell.value)
    assert isinstance(workbook["Recebimentos"]["A3"].value, datetime), "Data do recebimento não é célula Date"
    recebimentos = workbook["Recebimentos"]
    for row_index in range(3, 51):
        for column_index in (6, 7, 8, 10, 11, 12, 13, 14):
            cell = recebimentos.cell(row_index, column_index)
            assert isinstance(cell.value, (int, float)), f"Valor financeiro não numérico em {cell.coordinate}"
    all_visible_text = "\n".join(visible_text)
    assert not re.search(r"\b(CPF|PRONTU[ÁA]RIO|DIAGN[ÓO]STICO|UUID)\b", all_visible_text, re.I)
    workbook.close()
    return len(xml_entries), len(xml)


if __name__ == "__main__":
    pages, pdf_chars = verify_pdf()
    xml_entries, xml_chars = verify_xlsx()
    print({"pdf_pages": pages, "pdf_text_chars": pdf_chars, "xlsx_xml_entries": xml_entries, "xlsx_xml_chars": xml_chars})
